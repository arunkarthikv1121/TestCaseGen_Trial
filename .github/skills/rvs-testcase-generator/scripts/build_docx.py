#!/usr/bin/env python3
"""
Validate testcases.json and build the RV&S Word import file
(+ optional companion Excel sheet for the manually entered fields).

Usage (from the repository root):
    python .github/skills/rvs-testcase-generator/scripts/build_docx.py <run_folder>

    <run_folder> = output/<YYYY-MM-DD_HHMM>_<name>/  containing testcases.json

Writes into <run_folder>:
    <name>_RVS_import.docx       one heading per test case + step table (#, Actions, Expected Results)
    <name>_manual_fields.xlsx    ShortDescription, Preconditions, origin, source TC (if enabled)

Exit code 1 = validation errors (nothing written). Fix testcases.json and run again.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[4]
CONFIG = ROOT / "config" / "agent-config.yaml"
ORIGINS = {"NEW", "ADAPTED", "REUSED"}
TYPES = {"positive", "negative", "edge"}


def load_config() -> dict:
    with open(CONFIG, encoding="utf-8") as f:
        return yaml.safe_load(f)


# ---------------------------------------------------------------- validation
def validate(data: dict, cfg: dict) -> list[str]:
    errs: list[str] = []
    gen = cfg["generation"]
    tcs = data.get("test_cases")
    if not isinstance(tcs, list) or not tcs:
        return ["'test_cases' must be a non-empty list"]
    seen: set[str] = set()
    for i, tc in enumerate(tcs, 1):
        where = f"test_cases[{i}] '{tc.get('title', '?')}'"
        title = (tc.get("title") or "").strip()
        if not title:
            errs.append(f"{where}: title missing")
        elif title.lower() in seen:
            errs.append(f"{where}: duplicate title")
        seen.add(title.lower())
        if tc.get("origin") not in ORIGINS:
            errs.append(f"{where}: origin must be one of {sorted(ORIGINS)}")
        if tc.get("type") not in TYPES:
            errs.append(f"{where}: type must be one of {sorted(TYPES)}")
        if tc.get("origin") in {"ADAPTED", "REUSED"} and not tc.get("source_tc"):
            errs.append(f"{where}: source_tc required for {tc.get('origin')}")
        if not tc.get("requirement_refs"):
            errs.append(f"{where}: requirement_refs missing (use requirement heading/ID or file name)")
        steps = tc.get("steps") or []
        if tc.get("origin") != "REUSED":
            if not (gen["min_steps_per_testcase"] <= len(steps) <= gen["max_steps_per_testcase"]):
                errs.append(
                    f"{where}: {len(steps)} steps, allowed "
                    f"{gen['min_steps_per_testcase']}-{gen['max_steps_per_testcase']}"
                )
            for j, s in enumerate(steps, 1):
                if not (s.get("action") or "").strip():
                    errs.append(f"{where} step {j}: action empty")
                if not (s.get("expected") or "").strip():
                    errs.append(f"{where} step {j}: expected result empty (use '-' only if truly none)")
    return errs


# ---------------------------------------------------------------- docx helpers
def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def new_document(cfg: dict) -> Document:
    template = ROOT / cfg["paths"]["templates"] / "rvs_template.docx"
    if template.exists():
        doc = Document(template)
        body = doc.element.body
        for child in list(body):  # keep styles, drop template content
            if child.tag.split("}")[-1] != "sectPr":
                body.remove(child)
        return doc
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    return doc


def add_step_table(doc, steps: list[dict], cfg: dict) -> None:
    w = cfg["rvs_word_import"]
    cols = w["step_table_columns"]
    table = doc.add_table(rows=1, cols=len(cols))
    try:
        table.style = w["table_style"]
    except Exception:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0]
    repeat_header(hdr)
    for cell, name in zip(hdr.cells, cols):
        cell.text = ""
        run = cell.paragraphs[0].add_run(name)
        run.bold = True
        shade(cell, w["header_shading_hex"])
    for n, s in enumerate(steps, 1):
        row = table.add_row().cells
        row[0].text = str(n)
        row[1].text = s["action"].strip()
        row[2].text = s["expected"].strip()
    # column widths: narrow '#', equal Actions / Expected (Word needs grid + cell widths)
    table.autofit = False
    widths = [Cm(1.0)] + [Cm(8.0)] * (len(cols) - 1)
    for col, wdt in zip(table.columns, widths):
        col.width = wdt
    for row in table.rows:
        for cell, wdt in zip(row.cells, widths):
            cell.width = wdt


# ---------------------------------------------------------------- builders
def build_docx(data: dict, cfg: dict, out: Path) -> int:
    w = cfg["rvs_word_import"]
    doc = new_document(cfg)
    tcs = [t for t in data["test_cases"] if w["include_reused_in_docx"] or t["origin"] != "REUSED"]
    group_by = w.get("group_by", "none")
    current_group = None
    for tc in tcs:
        if group_by != "none":
            group = tc.get("group") or (tc.get("requirement_refs") or ["General"])[0]
            if group != current_group:
                doc.add_heading(group, level=w["group_heading_level"])
                current_group = group
        doc.add_heading(tc["title"].strip(), level=w["testcase_heading_level"])
        add_step_table(doc, tc["steps"], cfg)
        doc.add_paragraph("")
    doc.save(out)
    return len(tcs)


def build_sheet(data: dict, out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Manual fields"
    head = ["Title (Text heading)", "ShortDescription", "Preconditions", "Type",
            "Origin", "Source TC", "Requirement refs", "Open points"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    for tc in data["test_cases"]:
        ws.append([
            tc["title"],
            tc.get("short_description", ""),
            "\n".join(f"- {p}" for p in tc.get("preconditions", [])),
            tc.get("type", ""),
            tc.get("origin", ""),
            tc.get("source_tc", ""),
            ", ".join(tc.get("requirement_refs", [])),
            "\n".join(tc.get("open_points", [])),
        ])
    widths = [45, 35, 45, 10, 10, 18, 25, 40]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(out)


def main() -> int:
    for stream in (sys.stdout, sys.stderr):  # Windows consoles: never crash on special characters
        try:
            stream.reconfigure(errors="replace")
        except AttributeError:
            pass
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    run = Path(sys.argv[1])
    run = run if run.is_absolute() else ROOT / run
    src = run / "testcases.json"
    if not src.exists():
        print(f"Not found: {src}", file=sys.stderr)
        return 2
    cfg = load_config()
    data = json.loads(src.read_text(encoding="utf-8"))
    errs = validate(data, cfg)
    if errs:
        print("VALIDATION FAILED - fix testcases.json:", file=sys.stderr)
        for e in errs:
            print("  -", e, file=sys.stderr)
        return 1
    name = data.get("name") or run.name
    docx_path = run / f"{name}_RVS_import.docx"
    n = build_docx(data, cfg, docx_path)
    print(f"Word import file : {docx_path.relative_to(ROOT)}  ({n} test cases)")
    if cfg.get("companion_sheet", {}).get("enabled", True):
        xlsx_path = run / f"{name}_manual_fields.xlsx"
        build_sheet(data, xlsx_path)
        print(f"Manual fields    : {xlsx_path.relative_to(ROOT)}")
    counts: dict[str, int] = {}
    for t in data["test_cases"]:
        counts[t["origin"]] = counts.get(t["origin"], 0) + 1
    print("Origins          : " + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())))
    tbd = sum(
        json.dumps(t).count(cfg["generation"]["placeholder_for_missing_values"])
        for t in data["test_cases"]
    )
    if tbd:
        print(f"Placeholders     : {tbd} x {cfg['generation']['placeholder_for_missing_values']} - review before import")
    return 0


if __name__ == "__main__":
    sys.exit(main())
