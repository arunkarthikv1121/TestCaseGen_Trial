#!/usr/bin/env python3
"""
Extract the content of requirement and test case files into plain Markdown
so Copilot can read them reliably.

Usage (from the repository root):
    python .github/skills/rvs-testcase-generator/scripts/read_inputs.py

Reads   : paths in config/agent-config.yaml (input/ and, if enabled, examples/)
Writes  : output/_work/requirements/*.md
          output/_work/existing_testcases/*.md
          output/_work/tc_library.json   (existing test cases parsed into title + steps)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
CONFIG = ROOT / "config" / "agent-config.yaml"
SUPPORTED = {".docx", ".xlsx", ".xlsm", ".md", ".txt", ".csv"}


def load_config() -> dict:
    with open(CONFIG, encoding="utf-8") as f:
        return yaml.safe_load(f)


def iter_block_items(doc):
    """Yield paragraphs and tables in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def cell_text(cell) -> str:
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()).replace("|", "/")


def table_to_md(table: Table) -> list[str]:
    rows = [[cell_text(c) for c in r.cells] for r in table.rows]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return out


def heading_level(p: Paragraph) -> int | None:
    name = (p.style.name if p.style is not None else "") or ""
    for prefix in ("Heading ", "Überschrift "):
        if name.startswith(prefix):
            try:
                return int(name[len(prefix):].strip())
            except ValueError:
                return None
    return None


def docx_to_md(path: Path) -> tuple[str, list[dict]]:
    """Return markdown text and a list of parsed test cases (heading + step table)."""
    doc = Document(path)
    lines: list[str] = []
    testcases: list[dict] = []
    current: dict | None = None
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            lvl = heading_level(block)
            if lvl:
                lines.append("")
                lines.append("#" * min(lvl, 6) + " " + text)
                current = {"title": text, "heading_level": lvl, "text": [], "steps": []}
                testcases.append(current)
            else:
                lines.append(text)
                if current is not None:
                    current["text"].append(text)
        else:
            lines.append("")
            lines += table_to_md(block)
            lines.append("")
            if current is not None and len(block.rows) > 1:
                header = [cell_text(c).lower() for c in block.rows[0].cells]
                if any("action" in h for h in header) or any("expect" in h for h in header):
                    a_idx = next((i for i, h in enumerate(header) if "action" in h), 1)
                    e_idx = next((i for i, h in enumerate(header) if "expect" in h), 2)
                    for r in block.rows[1:]:
                        cells = [cell_text(c) for c in r.cells]
                        if len(cells) > max(a_idx, e_idx):
                            current["steps"].append(
                                {"action": cells[a_idx], "expected": cells[e_idx]}
                            )
    # keep only headings that actually carry a step table
    testcases = [t for t in testcases if t["steps"]]
    return "\n".join(lines).strip() + "\n", testcases


def xlsx_to_md(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        rows = [
            ["" if v is None else str(v).replace("\n", " ").replace("|", "/") for v in row]
            for row in ws.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        out.append(f"## Sheet: {ws.title}")
        out.append("| " + " | ".join(rows[0]) + " |")
        out.append("|" + "---|" * width)
        out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        out.append("")
    return "\n".join(out)


def convert_folder(src: Path, dst: Path, tag: str, library: list[dict]) -> int:
    if not src.exists():
        return 0
    dst.mkdir(parents=True, exist_ok=True)
    count = 0
    for f in sorted(src.rglob("*")):
        if not f.is_file() or f.suffix.lower() not in SUPPORTED or f.name.startswith("~$"):
            continue
        try:
            if f.suffix.lower() == ".docx":
                md, tcs = docx_to_md(f)
                for t in tcs:
                    t["source_file"] = str(f.relative_to(ROOT))
                    t["source_kind"] = tag
                library.extend(tcs if "testcase" in tag else [])
            elif f.suffix.lower() in {".xlsx", ".xlsm"}:
                md = xlsx_to_md(f)
            else:
                md = f.read_text(encoding="utf-8", errors="replace")
        except Exception as exc:  # report and continue with the other files
            print(f"  ! could not read {f.name}: {exc}", file=sys.stderr)
            continue
        out = dst / (f.stem + ".md")
        out.write_text(f"<!-- source: {f.relative_to(ROOT)} -->\n" + md, encoding="utf-8")
        print(f"  {tag:<28} {f.relative_to(ROOT)}  ->  {out.relative_to(ROOT)}")
        count += 1
    return count


def main() -> int:
    for stream in (sys.stdout, sys.stderr):  # Windows consoles: never crash on special characters
        try:
            stream.reconfigure(errors="replace")
        except AttributeError:
            pass
    cfg = load_config()
    p = cfg["paths"]
    work = ROOT / p["work"]
    library: list[dict] = []
    print("Extracting inputs ...")
    n_req = convert_folder(ROOT / p["requirements"], work / "requirements", "requirement", library)
    n_tc = convert_folder(ROOT / p["existing_testcases"], work / "existing_testcases", "existing_testcase", library)
    ex = cfg.get("examples", {})
    if ex.get("use_examples"):
        convert_folder(ROOT / ex["requirements"], work / "examples" / "requirements", "example_requirement", library)
        convert_folder(ROOT / ex["existing_testcases"], work / "examples" / "existing_testcases", "example_testcase", library)
    (work).mkdir(parents=True, exist_ok=True)
    (work / "tc_library.json").write_text(json.dumps(library, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nRequirement files : {n_req}")
    print(f"Existing TC files : {n_tc}")
    print(f"Parsed test cases : {len(library)}  ->  {(work / 'tc_library.json').relative_to(ROOT)}")
    if n_req == 0:
        print("\nNo requirement files found in", p["requirements"], file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
